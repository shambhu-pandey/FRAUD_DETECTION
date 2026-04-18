from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

TEMPLATE = r"c:\Users\REALME\OneDrive\Desktop\FRAUD_DETECTION\PROJECT WORK REVIEW 2 PPT TEMPLATE.pptx"
OUTPUT = r"c:\Users\REALME\OneDrive\Desktop\FRAUD_DETECTION\Shambhu_Pandey_Project_from_template.pptx"
STUDENT_NAME = "Shambhu Pandey"
REG_NO = "25MCA1020"
GUIDE_NAME = "Tamilrashi"

prs = Presentation(TEMPLATE)

# Add footer/contact box on first slide
if prs.slides:
    slide = prs.slides[0]
    left = Inches(0.5)
    top = Inches(6.5)
    width = Inches(9)
    height = Inches(0.6)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = f"Name: {STUDENT_NAME}    Reg No: {REG_NO}    Guide: {GUIDE_NAME}"
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)

# Add a demo screenshot placeholder slide at the end
blank_layout = None
# try to find a blank layout
for layout in prs.slide_layouts:
    if not layout.placeholders:
        blank_layout = layout
        break
if blank_layout is None:
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]

slide = prs.slides.add_slide(blank_layout)
# Title
title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9.0), Inches(0.6))
tf = title_box.text_frame
tf.text = "Demo — Project Screenshot"
for paragraph in tf.paragraphs:
    for run in paragraph.runs:
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0, 0, 0)

# Screenshot placeholder box
ph = slide.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(8.0), Inches(4.8))
tf2 = ph.text_frame
tf2.text = "[PASTE DEMO PROJECT SCREENSHOT HERE]"
for paragraph in tf2.paragraphs:
    for run in paragraph.runs:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(128, 128, 128)

prs.save(OUTPUT)
print(f"Saved modified presentation to: {OUTPUT}")
