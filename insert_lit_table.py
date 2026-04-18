from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

INPUT = r"c:\Users\REALME\OneDrive\Desktop\FRAUD_DETECTION\PROJECT_WORK_REVIEW_2_PPT_TEMPLATE_filled.pptx"
OUTPUT = r"c:\Users\REALME\OneDrive\Desktop\FRAUD_DETECTION\PROJECT_WORK_REVIEW_2_PPT_TEMPLATE_lit_filled.pptx"

prs = Presentation(INPUT)

# Find slide titled '3. LITERATURE SURVEY' (case-insensitive)
target_slide = None
for slide in prs.slides:
    if slide.shapes.title and 'LITERATURE SURVEY' in slide.shapes.title.text.upper():
        target_slide = slide
        break

if target_slide is None:
    # fallback: use slide index 2 (third slide)
    target_slide = prs.slides[2]

# Remove existing body placeholder text if present
for shape in list(target_slide.shapes):
    if not shape.has_table and shape.has_text_frame and shape is not target_slide.shapes.title:
        try:
            target_slide.shapes._spTree.remove(shape._element)
        except Exception:
            pass

# Define literature entries
rows = 11  # header + 10 entries
cols = 5
entries = [
    ("1","Statistical Fraud Detection Review","Statistical tests, rule-based","Surveys classical approaches and common fraud patterns","Older methods struggle with evolving fraud"),
    ("2","Data Mining for Financial Fraud","Decision trees, SVM, ensembles","ML improves detection over rules","Class imbalance reduces effectiveness"),
    ("3","XGBoost for Classification","Gradient boosting (XGBoost)","High accuracy and strong performance on imbalanced data","Requires careful hyperparameter tuning"),
    ("4","SHAP: Model Explanations","SHAP value explanations","Provides clear feature contributions","Computational cost for large datasets"),
    ("5","SMOTE: Oversampling Minority","SMOTE resampling + classifiers","Improves recall for minority class","Risk of overfitting synthetic samples"),
    ("6","Isolation Forest for Anomalies","Isolation Forest (unsupervised)","Detects novel/rare fraud patterns","High false positive rate on noisy data"),
    ("7","Autoencoder Anomaly Detection","Autoencoders (neural nets)","Finds unseen/novel frauds via reconstruction","Needs large clean data and tuning"),
    ("8","Cost‑Sensitive Learning","Cost-sensitive loss, class weighting","Reduces business cost by penalizing FN/FP","Choosing costs requires domain knowledge"),
    ("9","Real‑time Fraud Systems","Stream processing + lightweight ML","Supports low-latency scoring for live transactions","Engineering complexity and deployment overhead"),
    ("10","Ensemble & Hybrid Approaches","Stacking/ensembles (LR, RF, XGBoost)","Ensembles improve robustness and metrics","Increased inference latency and complexity"),
]

# Table placement and size
left = Inches(0.2)
top = Inches(1.3)
width = Inches(9.6)
height = Inches(5.5)

table_shape = target_slide.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

# Set column widths (approx)
col_widths = [Inches(0.5), Inches(3.0), Inches(2.2), Inches(2.4), Inches(1.5)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

# Header
headers = ["S.No", "Article Title", "Techniques Used", "Key Findings", "Limitations"]
for c, h in enumerate(headers):
    cell = table.cell(0, c)
    cell.text = h
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.size = Pt(12)

# Fill rows
for r, entry in enumerate(entries, start=1):
    for c, val in enumerate(entry):
        cell = table.cell(r, c)
        cell.text = val
        # adjust font size
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(11)

prs.save(OUTPUT)
print(f"Saved updated PPT with literature slide to: {OUTPUT}")
