from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Project details (edit if needed)
PROJECT_TITLE = "Fraud Detection System"
STUDENT_NAME = "Shambhu Pandey"
REG_NO = "25MCA1020"
GUIDE_NAME = "Tamilrashi"
OUTPUT_FILE = "Shambhu_Pandey_Project.pptx"


def add_title_slide(prs, title, subtitle):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    try:
        slide.placeholders[1].text = subtitle
    except Exception:
        pass


def add_bullet_slide(prs, heading, bullets):
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = heading
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        if i == 0:
            body.text = b
        else:
            p = body.add_paragraph()
            p.text = b
            p.level = 0


def add_two_col_slide(prs, heading, left_points, right_points):
    layout = prs.slide_layouts[3] if len(prs.slide_layouts) > 3 else prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = heading
    # Attempt to place two text boxes side by side
    left = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(4.2), Inches(4.0))
    rleft = left.text_frame
    for i, p in enumerate(left_points):
        if i == 0:
            rleft.text = p
        else:
            rp = rleft.add_paragraph()
            rp.text = p
    right = slide.shapes.add_textbox(Inches(5.0), Inches(1.6), Inches(4.0), Inches(4.0))
    rright = right.text_frame
    for i, p in enumerate(right_points):
        if i == 0:
            rright.text = p
        else:
            rp = rright.add_paragraph()
            rp.text = p


def add_screenshot_placeholder_slide(prs, heading):
    # Blank-like layout
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9.0), Inches(0.8))
    title_tf = title_box.text_frame
    title_tf.text = heading
    # Add a placeholder textbox where user can paste demo screenshot later
    ph = slide.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(8.0), Inches(4.8))
    tf = ph.text_frame
    tf.text = "DEMO PROJECT SCREENSHOT (paste image here)"
    # style hint
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(128, 128, 128)


def build_presentation(out_path=OUTPUT_FILE):
    prs = Presentation()
    # Title slide
    subtitle = f"Name: {STUDENT_NAME}    Reg No: {REG_NO}    Guide: {GUIDE_NAME}"
    add_title_slide(prs, PROJECT_TITLE, subtitle)

    # Objective
    add_bullet_slide(prs, "Objective", [
        "Build an effective fraud detection system",
        "Preprocess transactional datasets",
        "Train and evaluate classification models",
        "Select thresholds for production usage"
    ])

    # Dataset
    add_two_col_slide(prs, "Dataset", [
        "Sources: PaySim, CICIDS, IEEE, BankSim",
        "Rows: varied per dataset",
        "Key features: amount, time, source/dest"
    ], [
        "Cleaning: remove duplicates, format columns",
        "Feature engineering: aggregation, encoding",
        "Train/test split and resampling"
    ])

    # Methodology
    add_bullet_slide(prs, "Methodology", [
        "Exploratory data analysis and correlation checks",
        "Feature selection and imbalance handling (SMOTE/Tomek)",
        "Modeling: tree-based ensembles, logistic regression",
        "Threshold tuning and evaluation using selected metrics"
    ])

    # Results / Metrics
    add_bullet_slide(prs, "Key Results", [
        "Evaluation metrics computed and saved in models/",
        "Selected thresholds in models/thresholds_selected.json",
        "Performance: precision, recall, F1 (see models/evaluation_metrics.json)"
    ])

    # Demo placeholder slide
    add_screenshot_placeholder_slide(prs, "Demo — Project Screenshot")

    # Conclusions
    add_bullet_slide(prs, "Conclusions", [
        "Model achieves competitive detection performance",
        "Threshold tuning critical for production trade-offs",
        "Future work: online learning, feature drift monitoring"
    ])

    # Acknowledgement / Contact
    add_bullet_slide(prs, "Acknowledgements & Contact", [
        f"Student: {STUDENT_NAME}",
        f"Guide: {GUIDE_NAME}",
        "Questions welcome"
    ])

    prs.save(out_path)
    print(f"Saved presentation to {out_path}")


if __name__ == '__main__':
    build_presentation()
